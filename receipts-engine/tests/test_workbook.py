"""Workbook fill writes only yellow input cells and never touches formulas (BUILD_SPEC 9.6).

The real workbook is human-provided (input/), so this test builds a stand-in with the documented layout.
"""
from __future__ import annotations

import csv
import shutil

import pytest
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from receipts.models import LEAD_COLUMNS
from receipts.workbook import WorkbookMappingError, fill_workbook
from tests.conftest import ROOT

YELLOW = PatternFill("solid", fgColor="FFFFFF00")


def _yellow(ws, rng):
    for row in ws[rng]:
        for c in row:
            c.fill = YELLOW


def make_workbook(path):
    wb = Workbook()
    wb.remove(wb.active)
    st = wb.create_sheet("00_START")
    for r in range(5, 19):
        st[f"A{r}"] = f"Label {r}"
    _yellow(st, "B5:B18")
    st["A21"] = "I accept the declaration"
    _yellow(st, "B21:B21")
    _yellow(st, "F5:F10")
    st["G5"] = '=COUNTIF(F5:F10,"Yes")'
    ins = wb.create_sheet("01_INSIGHT")
    _yellow(ins, "B6:F10")
    _yellow(ins, "B14:F16")
    ins["A20"] = "=COUNTA(B6:F10)"
    leads = wb.create_sheet("02_LEADS")
    for j, name in enumerate(LEAD_COLUMNS):
        leads.cell(row=4, column=j + 1, value=name)
    _yellow(leads, "A5:AB1004")
    for r in range(5, 1005):
        leads[f"AC{r}"] = f'=IF(COUNTIF($H$5:$H$1004,H{r})>1,"DUPLICATE","UNIQUE")'
        leads[f"AD{r}"] = f'=IF(COUNTA(A{r}:AB{r})>=27,"COMPLETE","INCOMPLETE")'
    eng = wb.create_sheet("03_ENGINE_LAUNCH")
    _yellow(eng, "B6:C14")
    _yellow(eng, "H6:H14")
    _yellow(eng, "B18:H21")
    _yellow(eng, "B24:B31")
    eng["A40"] = '=COUNTIF(B6:B14,"Yes")'
    sc = wb.create_sheet("04_SELF_CHECK")
    _yellow(sc, "C6:C12")
    _yellow(sc, "E6:E12")
    _yellow(sc, "C24:C28")
    _yellow(sc, "E24:E28")
    sc["D30"] = "=SUM(C24:C28)"
    wb.save(path)


def fake_leads_csv(path, n):
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(LEAD_COLUMNS)
        for i in range(n, 0, -1):  # unsorted on purpose
            row = {c: f"x{i}" for c in LEAD_COLUMNS}
            row.update(lead_id=f"L{i:04d}", content_count_30d="9", longform_count_60d="3", lead_score="80.5",
                       followers_or_subscribers="12000", rejection_reason="")
            w.writerow([row[c] for c in LEAD_COLUMNS])


@pytest.fixture
def human_cfg(tmp_path):
    d = tmp_path / "cfg"
    shutil.copytree(ROOT / "config", d)
    sub = yaml.safe_load((d / "submission.yaml").read_text())
    sub["start"] = {k: f"value {k}" for k in sub["start"]}
    sub["accept_declaration"] = True
    proof = tmp_path / "proof.txt"
    proof.write_text("ok")
    sub["checklist"] = {"F5": str(proof), "F6": "https://example.test/video", "F7": "", "F8": str(proof),
                        "F9": "missing/file.md", "F10": str(proof)}
    sub["engine_requirements"] = {r: {"evidence": str(proof), "note": f"note {r}"} for r in range(6, 15)}
    sub["launch_kit_path"] = str(tmp_path / "launch_kit.md")
    (tmp_path / "launch_kit.md").write_text("# Kit\n## 1) Who\nCoaches first.\n## 2) Channels\nEmail first.\n")
    (d / "submission.yaml").write_text(yaml.safe_dump(sub))
    (d / "insight.yaml").write_text(yaml.safe_dump({"observations": [[f"o{i}{j}" for j in range(5)] for i in range(5)],
                                                    "opportunities": [[f"p{i}{j}" for j in range(5)] for i in range(3)]}))
    (d / "selfscore.yaml").write_text(yaml.safe_dump({"scores": {r: {"score": 3, "reason": f"r{r}"} for r in range(24, 29)}}))
    return d


def snapshot_cells(path):
    wb = load_workbook(path)
    return {(ws.title, c.coordinate): c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row}


def test_fill_writes_only_yellow_cells_and_keeps_formulas(settings, db, tmp_path, human_cfg):
    src, dst, leads = tmp_path / "wb.xlsx", tmp_path / "out" / "workbook_fill.xlsx", tmp_path / "leads.csv"
    make_workbook(src)
    fake_leads_csv(leads, 1200)
    before = snapshot_cells(src)
    fill_workbook(db, settings, src, dst, leads, human_cfg)
    after = snapshot_cells(dst)
    wb_after = load_workbook(dst)
    formulas_before = {k: v for k, v in before.items() if isinstance(v, str) and v.startswith("=")}
    assert formulas_before and all(after[k] == v for k, v in formulas_before.items())
    changed = [k for k in after if after.get(k) != before.get(k)]
    assert changed
    for sheet, coord in changed:
        fill = wb_after[sheet][coord].fill.fgColor.rgb
        assert fill == "FFFFFF00", (sheet, coord)
    # header untouched, exactly 1,000 rows written in lead_id order, typed numbers
    lw = wb_after["02_LEADS"]
    assert [lw.cell(row=4, column=j + 1).value for j in range(30)] == LEAD_COLUMNS
    written = [r for r in range(5, 1005) if lw[f"A{r}"].value]
    assert len(written) == 1000 and lw["A5"].value == "L0001" and lw["A1004"].value == "L1000"
    assert lw["L5"].value == 9 and lw["Z5"].value == 80.5 and lw["AB5"].value in ("", None)
    assert "1200 rows" in wb_after["04_SELF_CHECK"]["E6"].value
    st = wb_after["00_START"]
    assert st["B21"].value == "Yes" and st["F5"].value == "Yes" and st["F7"].value == "No" and st["F9"].value == "No"
    assert wb_after["03_ENGINE_LAUNCH"]["B6"].value == "Yes"
    assert wb_after["03_ENGINE_LAUNCH"]["B24"].value == "Coaches first."


def test_fill_aborts_on_non_yellow_target(settings, db, tmp_path, human_cfg):
    src, dst, leads = tmp_path / "wb.xlsx", tmp_path / "out.xlsx", tmp_path / "leads.csv"
    make_workbook(src)
    wb = load_workbook(src)
    wb["00_START"]["B7"].fill = PatternFill()  # an input cell lost its yellow
    wb.save(src)
    fake_leads_csv(leads, 3)
    with pytest.raises(WorkbookMappingError) as e:
        fill_workbook(db, settings, src, dst, leads, human_cfg)
    assert "00_START!B7" in str(e.value) and not dst.exists()


def test_fill_aborts_when_target_holds_formula(settings, db, tmp_path, human_cfg):
    src, dst, leads = tmp_path / "wb.xlsx", tmp_path / "out.xlsx", tmp_path / "leads.csv"
    make_workbook(src)
    wb = load_workbook(src)
    wb["02_LEADS"]["C5"] = "=A5"
    wb.save(src)
    fake_leads_csv(leads, 3)
    with pytest.raises(WorkbookMappingError) as e:
        fill_workbook(db, settings, src, dst, leads, human_cfg)
    assert "02_LEADS!C5: holds a formula" in str(e.value)
